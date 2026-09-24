#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _prepend_block(doc: Document, paragraphs: list) -> None:
    body = doc._body._element
    first = body[0] if len(body) else None
    for p in reversed(paragraphs):
        if first is None:
            body.append(p._p)
        else:
            body.insert(body.index(first), p._p)


def _add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    added = []

    title = doc.add_paragraph()
    added.append(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("RLAIF Before vs After")
    r.bold = True
    r.font.size = Pt(28)

    subtitle = doc.add_paragraph("Executive Summary for Client Delivery")
    added.append(subtitle)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(13)

    meta = doc.add_paragraph(
        "Scope: Oldest repository baseline vs latest repository commit with commit-precise source traceability"
    )
    added.append(meta)
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in meta.runs:
        run.font.size = Pt(10)
        run.italic = True

    spacer1 = doc.add_paragraph()
    added.append(spacer1)

    h = doc.add_paragraph("Executive Highlights")
    added.append(h)
    hr = h.runs[0]
    hr.bold = True
    hr.font.size = Pt(15)

    bullets = [
        "Comparison is fully aligned to current repository history; no synthetic v1-v7 labeling is used.",
        "Each before/after snapshot is tied to exact commit SHAs, files, and line-focused code regions.",
        "Core deltas cover prompt/replay behavior, tool metadata, RCA controls, persistence, workflow gates, and judging logic.",
        "All image cards remain clickable and resolve to commit-precise GitHub source permalinks.",
        "Preflight and smoke validation were executed before final document publication.",
    ]
    for item in bullets:
        p = doc.add_paragraph(item, style="List Bullet")
        added.append(p)
        for run in p.runs:
            run.font.size = Pt(11)

    spacer2 = doc.add_paragraph()
    added.append(spacer2)

    note = doc.add_paragraph(
        "The following pages preserve the complete technical before/after evidence package without modification."
    )
    added.append(note)
    for run in note.runs:
        run.font.size = Pt(10)
        run.italic = True

    pb = doc.add_paragraph()
    added.append(pb)
    pb.add_run().add_break(WD_BREAK.PAGE)

    _prepend_block(doc, added)


def main() -> None:
    base_path = Path("output/rlaif_before_after.docx")
    out_path = Path("output/rlaif_before_after_exec.docx")

    if not base_path.exists():
        raise FileNotFoundError(f"Base document not found: {base_path}")

    out_doc = Document(str(base_path))
    _add_cover(out_doc)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_doc.save(str(out_path))

    print(str(out_path))


if __name__ == "__main__":
    main()
